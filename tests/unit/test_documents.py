"""Document rendering helpers and parameter validation."""

import base64
import json
import re
from io import BytesIO
from pathlib import Path
from typing import TYPE_CHECKING, Any

import pytest
from docx import Document
from PIL import Image
from pypdf import PdfReader

from jcpy.config.loader import build_config
from jcpy.documents import docx as docx_module
from jcpy.documents.docx import html_to_docx, sanitize_html
from jcpy.documents.pdf import GuardedFetcher, html_to_pdf, page_style
from jcpy.documents.resources import (
    ResourcePolicy,
    ResourceRefusedError,
    load_remote,
)
from jcpy.helpers import ssrf
from jcpy.v1.generate_docx.handler import validate as docx_issues
from jcpy.v1.generate_pdf.handler import validate as pdf_issues

if TYPE_CHECKING:
    from collections.abc import Callable

    from jcpy.types import JsonObject
    from jcpy.validation import Issue

CASES = json.loads(
    (
        Path(__file__).parent.parent
        / "fixtures"
        / "document_schema_cases.json"
    ).read_text()
)
VALIDATORS: dict[str, Callable[[JsonObject], list[Issue]]] = {
    "generatePdf": pdf_issues,
    "generateDocx": docx_issues,
}
OPEN = ResourcePolicy(remote=True, guard=True, limit=None, timeout=5)
CLOSED = ResourcePolicy(remote=False, guard=True, limit=None, timeout=5)


@pytest.mark.parametrize(
    "case",
    CASES,
    ids=[f"{c['schema']}|{json.dumps(c['input'])}" for c in CASES],
)
def test_validation_matches_zod(case: dict[str, Any]) -> None:
    issues = VALIDATORS[case["schema"]](case["input"])

    assert [(issue.path, issue.message) for issue in issues] == [
        (issue["path"], issue["message"]) for issue in case["issues"]
    ]


def test_policy_from_config() -> None:
    policy = ResourcePolicy.from_config(
        build_config(
            {
                "maxFileSize": "1kb",
                "timeoutLimit": 7,
                "allowPrivateNetworkUploads": True,
                "pdf": {"isRemoteEnabled": False},
            }
        )
    )
    unlimited = ResourcePolicy.from_config(build_config({"maxFileSize": ""}))

    assert policy == ResourcePolicy(
        remote=False, guard=False, limit=1024, timeout=7
    )
    assert unlimited.limit is None


async def test_load_remote_refusals() -> None:
    with pytest.raises(ResourceRefusedError, match="disabled"):
        await load_remote("http://example.com/a.png", CLOSED)
    with pytest.raises(ResourceRefusedError, match="Only http"):
        await load_remote("ftp://example.com/a.png", OPEN)


@pytest.mark.parametrize(
    ("paper", "orientation", "size"),
    [
        ("A4", "portrait", "210mm 297mm"),
        ("A3", "landscape", "420mm 297mm"),
        ("Tabloid", "portrait", "11in 17in"),
    ],
)
def test_page_style(paper: str, orientation: str, size: str) -> None:
    assert page_style(paper, orientation) == (
        f"@page {{ size: {size}; margin: 1cm; }}"
    )


def media_box(pdf: bytes) -> tuple[float, float]:
    box = PdfReader(BytesIO(pdf)).pages[0].mediabox
    return round(float(box.width), 1), round(float(box.height), 1)


async def test_pdf_page_size() -> None:
    portrait = await html_to_pdf("<p>x</p>", OPEN)
    landscape = await html_to_pdf("<p>x</p>", OPEN, "Letter", "landscape")

    assert portrait.startswith(b"%PDF-")
    assert media_box(portrait) == (595.3, 841.9)
    assert media_box(landscape) == (792.0, 612.0)


def png(color: str = "red", fmt: str = "PNG") -> bytes:
    output = BytesIO()
    Image.new("RGB", (4, 4), color).save(output, fmt)
    return output.getvalue()


def test_fetcher_refuses_local_files(tmp_path: Path) -> None:
    secret = tmp_path / "secret.txt"
    secret.write_text("secret")
    fetcher = GuardedFetcher(OPEN)

    with pytest.raises(ValueError, match="disallowed protocol"):
        fetcher.fetch(secret.as_uri())
    assert fetcher.fetch("data:text/plain,hi").read() == b"hi"


class Recorder:
    """Stand-in for ``ssrf.fetch`` recording requested URLs."""

    def __init__(self, body: bytes, content_type: str | None) -> None:
        self.body = body
        self.content_type = content_type
        self.urls: list[str] = []

    async def __call__(self, url: str, **_: object) -> ssrf.Fetched:
        self.urls.append(url)
        return ssrf.Fetched(url, self.body, self.content_type)


async def test_pdf_resources_go_through_the_policy(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    recorder = Recorder(png(), "image/png")
    monkeypatch.setattr(ssrf, "fetch", recorder)
    html = '<img src="https://cdn.example/a.png"><img src="https://x/b.png">'

    with_images = await html_to_pdf(html, OPEN)
    await html_to_pdf(html, CLOSED)

    assert recorder.urls == ["https://cdn.example/a.png", "https://x/b.png"]
    assert b"/Subtype /Image" in with_images


async def test_pdf_resource_without_content_type(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(ssrf, "fetch", Recorder(png(), None))

    pdf = await html_to_pdf('<img src="https://cdn.example/a">', OPEN)

    assert b"/Subtype /Image" in pdf


def data_uri(contents: bytes, kind: str = "png") -> str:
    return f"data:image/{kind};base64," + base64.b64encode(contents).decode()


async def test_sanitize_html(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    recorder = Recorder(png(fmt="WEBP"), "image/webp")
    monkeypatch.setattr(ssrf, "fetch", recorder)
    html = (
        "<style>p{}</style><script>alert(1)</script>"
        '<link rel="stylesheet" href="/etc/passwd"><p>text</p>'
        f'<img alt="data" src="{data_uri(png())}">'
        f'<img alt="webp" src="{data_uri(png(fmt="WEBP"), "webp")}">'
        '<img alt="remote" src="https://cdn.example/a.webp">'
        f'<img alt="local" src="{tmp_path}/a.png">'
        '<img alt="file" src="file:///etc/passwd">'
        '<img alt="bad" src="data:image/png;base64,***">'
        '<img alt="text" src="data:text/plain;base64,aGk=">'
        f'<img alt="nobase64" src="data:image/svg+xml,{"<svg/>"}">'
        '<img alt="junk" src="data:image/png;base64,aGk=">'
        '<img alt="empty">'
    )

    safe = await sanitize_html(html, OPEN)

    assert "<style" not in safe
    assert "<script" not in safe
    assert "<link" not in safe
    kept = re.findall(r'<img alt="(\w+)" src="data:image/(\w+);base64,', safe)
    assert kept == [("data", "png"), ("webp", "png"), ("remote", "png")]
    assert recorder.urls == ["https://cdn.example/a.webp"]


async def test_failed_remote_images_are_dropped(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    async def refuse(url: str, **_: object) -> ssrf.Fetched:
        raise ssrf.DownloadTooLargeError

    monkeypatch.setattr(ssrf, "fetch", refuse)

    safe = await sanitize_html('<img src="https://cdn.example/a.png">', OPEN)

    assert "<img" not in safe


async def test_docx_document(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(ssrf, "fetch", Recorder(png(), "image/png"))
    html = (
        "<h1>Title</h1><p><b>bold</b> text</p><ul><li>item</li></ul>"
        "<table><tr><td>cell</td></tr></table>"
        '<img src="https://cdn.example/a.png">'
    )

    contents = await html_to_docx(html, OPEN)

    document = Document(BytesIO(contents))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Title" in text
    assert "bold text" in text
    assert "item" in text
    assert document.tables[0].cell(0, 0).text == "cell"
    assert len(document.inline_shapes) == 1
    section = document.sections[0]
    assert section.left_margin == docx_module.MARGIN
    assert section.top_margin == docx_module.MARGIN
